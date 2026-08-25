import os
import tempfile
import unittest

from openpyxl import Workbook

from autopptx2.datasource import (
    ExcelSource,
    build_merged_groups,
    compare_list_with_master,
    expected_photo_count,
)


class ListMasterCompareTests(unittest.TestCase):
    def test_compare_list_loader_combines_list_sheets_only(self):
        wb = Workbook()
        university = wb.active
        university.title = 'University List'
        university.append(['Code_RP', 'Name'])
        university.append(['UNI01', 'University A'])
        building = wb.create_sheet('Building List')
        building.append(['Report Code', 'Name of Block'])
        building.append(['BLD01', 'Building A'])
        summary = wb.create_sheet('Summary')
        summary.append(['Code_RP', 'Name'])
        summary.append(['SHOULD_SKIP', 'Summary row'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(path)
            src = ExcelSource()
            self.assertTrue(src.load_compare_list(path), src.error)
            self.assertEqual(
                [(row['Code_RP'], row['Name'], row['_SourceSheet']) for row in src.rows],
                [
                    ('UNI01', 'University A', 'University List'),
                    ('BLD01', 'Building A', 'Building List'),
                ],
            )
        finally:
            wb.close()
            if os.path.exists(path):
                os.remove(path)

    def test_expected_photo_count_uses_screen_quantity_and_safe_fallback(self):
        self.assertEqual(expected_photo_count({'Channel': 'University', 'DP': 2}), 2)
        self.assertEqual(
            expected_photo_count({'Channel': 'Coffee & Milk Tea', 'Quantity': 3}), 3)
        self.assertEqual(
            expected_photo_count({'Channel': 'Hotel & Resort', 'Quantity': 80}), 1)

    def test_compare_groups_missing_images_by_master_channel(self):
        master = {
            'UNI01': {'Code_RP': 'UNI01', 'Name': 'University A',
                      'Channel': 'University', 'DP': 2},
            'BLD01': {'Code_RP': 'BLD01', 'Name': 'Building A',
                      'Channel': 'Building', 'LCD': 1},
            'COF01': {'Code_RP': 'COF01', 'Name': 'Coffee A',
                      'Channel': 'Coffee & Milk Tea', 'Quantity': 3},
        }
        list_rows = [
            {'Code_RP': 'UNI01'},
            {'Code_RP': 'BLD01'},
            {'Code_RP': 'COF01'},
            {'Code_RP': 'UNKNOWN', 'Channel': 'Fastfood'},
            {'Code_RP': 'uni01'},  # duplicate must count once
        ]
        images = {
            'UNI01': ['uni-1.jpg'],
            'BLD01': ['bld-1.jpg'],
        }
        report = compare_list_with_master(
            list_rows, master, images, build_merged_groups(master))

        self.assertEqual(report['n_list'], 4)
        self.assertEqual(report['n_master_matches'], 3)
        self.assertEqual(report['n_not_in_master'], 1)
        self.assertEqual(report['n_ok_sites'], 1)
        self.assertEqual(report['n_missing_sites'], 3)
        self.assertEqual(report['n_missing_images'], 5)

        by_code = {row['code']: row for row in report['rows']}
        self.assertEqual(by_code['UNI01']['missing'], 1)
        self.assertEqual(by_code['BLD01']['status'], 'ok')
        self.assertEqual(by_code['COF01']['missing'], 3)
        self.assertEqual(by_code['UNKNOWN']['status'], 'not_in_master')

        by_channel = {row['channel']: row for row in report['channels']}
        self.assertEqual(by_channel['University']['missing'], 1)
        self.assertEqual(by_channel['Building']['ok_sites'], 1)
        self.assertEqual(by_channel['Coffee & Milk Tea']['missing'], 3)
        self.assertEqual(by_channel['Fastfood']['not_in_master'], 1)

    def test_compare_codes_without_images_does_not_report_false_missing(self):
        master = {
            'UNI01': {'Code_RP': 'UNI01', 'Name': 'University A',
                      'Channel': 'University', 'DP': 2},
            'BLD01': {'Code_RP': 'BLD01', 'Name': 'Building A',
                      'Channel': 'Building', 'LCD': 1},
        }
        list_rows = [
            {'Code_RP': 'UNI01'},
            {'Code_RP': 'BLD01'},
            {'Code_RP': 'UNKNOWN', 'Channel': 'Fastfood'},
        ]

        report = compare_list_with_master(
            list_rows, master, {}, build_merged_groups(master),
            check_images=False)

        self.assertFalse(report['images_checked'])
        self.assertEqual(report['n_list'], 3)
        self.assertEqual(report['n_master_matches'], 2)
        self.assertEqual(report['n_not_in_master'], 1)
        self.assertEqual(report['n_missing_sites'], 0)
        self.assertEqual(report['n_missing_images'], 0)
        self.assertEqual(report['n_required'], 0)
        by_code = {row['code']: row for row in report['rows']}
        self.assertEqual(by_code['UNI01']['status'], 'ok')
        self.assertEqual(by_code['BLD01']['status'], 'ok')
        self.assertEqual(by_code['UNKNOWN']['status'], 'not_in_master')


if __name__ == '__main__':
    unittest.main()
