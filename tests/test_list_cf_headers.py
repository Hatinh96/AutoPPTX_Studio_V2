import os
import tempfile
import unittest

from openpyxl import Workbook

from autopptx2.datasource import ExcelSource, screen_qty


def _list_cf(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Coffee Shop & Milk Tea List'
    ws.append(['COFFEE SHOP & MILK TEA LIST'])
    ws.append(['No.', 'BD Code', 'Chanel', 'Report Code', 'Name', 'Address', 'Ward',
               'Dist (old)', 'City', 'Region', 'Quantity of ', None, None, None,
               'Traffic (**)', None, 'Uptime', 'Note'])
    ws.append([None] * 10 + [' LCD', 'Digital Poster', 'Digital Standee ', 'Total',
                             'Day', 'Week'])
    ws.append([1, 'NGANBD', 'MT', 'PLBINHPHU', 'Phúc Long Bình Phú', '16 Bình Phú',
               'Bình Phú', 6, 'Hồ Chí Minh', 'Miền Nam', None, 4, None, 4, 1100, 7700,
               '7am-10pm', 'ON AIR'])
    ws.append([2, 'NGANBD', 'CF', 'HLCALMETTE', 'Highlands Coffee Calmette',
               '161-163 Calmette', 'Bến Thành', 1, 'Hồ Chí Minh', 'Miền Nam',
               1, 3, None, 4, 900, 6300, '7am-11pm', 'ON AIR'])
    ws.append([2, None, None, None, None, '  Total Ho Chi Minh', None, None, None, None,
               1, 7, 0, 8, 2000, 14000, None, None])
    ws.append([3, 'ANHBD', 'CF', 'HLSTARLAOCAI', 'Highlands Coffee Star Lào Cai',
               '03 Hoàng Liên', 'Lào Cai', None, 'Lào Cai', 'Miền Bắc',
               None, 2, None, 2, 500, 3500, '7am-11pm', 'ON AIR'])
    wb.save(path)


class TestListCfHeaders(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self._tmp.name, 'List CF.xlsx')
        _list_cf(self.path)
        src = ExcelSource()
        self.assertTrue(src.load(self.path), src.error)
        self.by_code = src.by_code(None, include_off=True)

    def tearDown(self):
        self._tmp.cleanup()

    def test_digital_poster_counts_as_dp(self):
        row = self.by_code['PLBINHPHU']
        self.assertEqual(row['DP'], 4)
        self.assertEqual(screen_qty(row), 4)

    def test_grouped_subheaders_use_parent(self):
        row = self.by_code['HLCALMETTE']
        self.assertEqual(row['Quantity'], 4)
        self.assertEqual(row['TrafficDay'], 900)
        self.assertEqual(row['TrafficWeek'], 6300)
        self.assertEqual(str(row['District']), '1')

    def test_province_total_row_not_merged_into_store(self):
        self.assertEqual(screen_qty(self.by_code['HLCALMETTE']), 4)
        self.assertEqual(sorted(self.by_code), ['HLCALMETTE', 'HLSTARLAOCAI', 'PLBINHPHU'])

    def test_district_not_carried_into_other_province(self):
        self.assertEqual(self.by_code['HLSTARLAOCAI'].get('District'), '')


def _hotel_building(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Hotel and Resort List'
    ws.append(['HOTEL AND RESORT LIST'])
    ws.append(['No', 'BD CODE', 'REPORT CODE', 'Type', 'Name', 'Address', 'Ward',
               'Dist (old)', 'City', 'Star', 'Q.ty of Elevator', 'Position Booking',
               None, None, None, None, None, 'Quantity of Rooms',
               'Traffic/day\n(Traffic =  Quantity of Room*3*70%)', 'Traffic/wk'])
    ws.append([None] * 11 + ['Digital Poster\n Inside Elevator',
                             'Digital Poster\n Front of Elevator', 'LCD\nFront of Elevator',
                             'Digital Standee',
                             'Digital Poster Others ( Buffet, Facilities...)', 'Total'])
    ws.append([None, None, 'HOTEL AND RESORT IN DA NANG'])
    ws.append([1, 'ANHBD', 'KOIRESORT', 'Resort', 'Koi Resort', '1 Võ Nguyên Giáp',
               'Mỹ An', 'Ngũ Hành Sơn', 'Đà Nẵng', 5, 4, 8, 4, None, 2, 6, 20,
               300, 630, 4410])
    bs = wb.create_sheet('Building List')
    bs.append(['BUILDING LIST'])
    bs.append(['No.', 'Type', 'BD Code', 'Report Code', 'Name of Block', 'Address (mới)',
               'Ward (mới)', 'City (mới)', 'Address (cũ)', 'District (cũ)', 'City (cũ)',
               'Traffic/week', 'Q.ty of Lift', 'Q.ty of GP \ninside', 'Qty of GP\n Ground',
               'Qty of GP \nFacilities floor\n(Tầng tiện ích: Gym, Thư viện, hồ bơi,..) ',
               'Qty of GP \nParking floor\n(Tầng gửi xe) ',
               'Qty of GP\n stay\n(Đã trừ các tầng G, tầng gửi xe trên mặt đất và tầng tiện ích) ',
               'Q.ty GP of AP'])
    bs.append([1, 'AP', 'PHUBD', 'CENTRALGARDENA', 'Central Garden Block A',
               '225 Bến Chương Dương', 'Cầu Ông Lãnh', 'Hồ Chí Minh',
               '225 Bến Chương Dương', 1, 'Hồ Chí Minh', 438, 3, 3, 3, 1, 3, 2, 12])
    wb.save(path)


class TestHotelBuildingHeaders(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        path = os.path.join(self._tmp.name, 'W38_Update Golden 2026.xlsx')
        _hotel_building(path)
        src = ExcelSource()
        self.assertTrue(src.load(path), src.error)
        self.by_code = src.by_code(None, include_off=True)

    def tearDown(self):
        self._tmp.cleanup()

    def test_area_label_in_code_column_is_not_a_site(self):
        self.assertEqual(sorted(self.by_code), ['CENTRALGARDENA', 'KOIRESORT'])

    def test_hotel_counts_every_digital_poster_column(self):
        row = self.by_code['KOIRESORT']
        self.assertEqual(screen_qty(row), 20)
        self.assertEqual(row['TrafficDay'], 630)

    def test_building_long_gp_headers(self):
        row = self.by_code['CENTRALGARDENA']
        self.assertEqual(screen_qty(row), 12)
        self.assertEqual(str(row['District']), '1')


def _audit_book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Coffee Shop & Milk Tea List'
    ws.append(['No.', 'Report Code', 'Name', 'Address', 'City', 'Quantity of ', None, None,
               'Screen Wall', 'Q.ty of Floor (For Giant Poster Avalable)'])
    ws.append([None] * 5 + ['LCD', 'Digital Poster', 'Total'])
    ws.append([1, 'CFA', 'Cafe A', '1 Lê Lợi', 'Hồ Chí Minh', 1, 3, 4, None, 12])
    ws.append([2, 'CFB', 'Cafe B', '2 Lê Lợi', 'Hồ Chí Minh', 1, 1, 5, 2, 9])
    ws.append([None, None, None, None, None, 2, 4, 9, 2, 21])
    bs = wb.create_sheet('Building List')
    bs.append(['No.', 'Report Code', 'Name of Block', 'Address', 'District', 'City',
               'Q.ty of Floor', 'Q.ty of Apartment', 'Traffic/week', 'Q.ty of GP inside',
               'Qty of GP Ground'])
    for i in range(6):
        bs.append([i + 1, f'AP{i}', f'Block {i}', f'{i} Nguyễn Trãi', 1, 'Hồ Chí Minh',
                   20, 300, 900, 2, 1])
    bs.append([7, 'APSHIFT', 'Block S', '9 Nguyễn Trãi', 'Bến Thành', 'Hồ Chí Minh',
               '9 Nguyễn Trãi, Phường 2', 'Quận 1', 'Hồ Chí Minh', 900, 34000])
    bs.append([8, 'APHUGE', 'Block H', '10 Nguyễn Trãi', 1, 'Hồ Chí Minh',
               20, 300, 900, 2, 64000])
    bs.append([9, 'APBIG', 'Block B', '11 Nguyễn Trãi', 1, 'Hồ Chí Minh',
               20, 300, 900, 2, 1010])
    wb.save(path)


class TestLoadAudit(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        path = os.path.join(self._tmp.name, 'W40 Golden 2026.xlsx')
        _audit_book(path)
        self.src = ExcelSource()
        self.assertTrue(self.src.load(path), self.src.error)
        self.by_code = self.src.by_code(None, include_off=True)
        self.text = '\n'.join(self.src.warnings)

    def tearDown(self):
        self._tmp.cleanup()

    def test_floor_count_for_giant_poster_is_not_gp(self):
        self.assertEqual(screen_qty(self.by_code['CFA']), 4)

    def test_unlabeled_total_row_is_skipped(self):
        self.assertEqual(screen_qty(self.by_code['CFB']), 2)

    def test_warns_when_sum_differs_from_file_total(self):
        self.assertIn('CFB', self.text)
        self.assertIn('khác cột Total', self.text)

    def test_warns_about_unknown_numeric_column(self):
        self.assertIn('«Screen Wall»', self.text)
        self.assertNotIn('Floor', self.text)

    def test_shifted_row_numbers_dropped_and_reported(self):
        self.assertEqual(screen_qty(self.by_code['APSHIFT']), 0)
        self.assertIn('lệch cột', self.text)
        self.assertEqual(screen_qty(self.by_code['AP0']), 3)

    def test_warns_about_impossible_screen_count(self):
        self.assertIn('APHUGE', self.text)
        self.assertIn('bất thường', self.text)
        self.assertIn('1 dòng có số màn bất thường', self.text)
