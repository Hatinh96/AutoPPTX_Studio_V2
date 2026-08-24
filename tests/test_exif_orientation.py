"""Portrait JPEG + EXIF Orientation 6/8 must stay upright (no extra 90°)."""
import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from PIL import Image

from autopptx2 import geometry as G
from autopptx2 import effects as FX
from autopptx2 import exporter as EX


def _marked_portrait(w=60, h=120):
    """Ảnh dọc: đỉnh đỏ, mép trái xanh, nền xanh dương."""
    im = Image.new('RGB', (w, h), (20, 40, 200))
    for y in range(min(24, h)):
        for x in range(w):
            im.putpixel((x, y), (220, 30, 30))
    for y in range(h):
        for x in range(min(10, w)):
            im.putpixel((x, y), (20, 200, 40))
    return im


def _save_orient(im, path, orientation):
    exif = Image.Exif()
    exif[G.EXIF_ORIENTATION] = orientation
    im.save(path, 'JPEG', quality=95, exif=exif)


def _phone_portrait_jpeg(path, orientation=6):
    """Giống máy ảnh: pixel ngang + tag 6/8, khi xem phải ra ảnh dọc."""
    visual = _marked_portrait()
    if orientation == 6:
        stored = visual.transpose(Image.Transpose.ROTATE_90)
    elif orientation == 8:
        stored = visual.transpose(Image.Transpose.ROTATE_270)
    else:
        stored = visual
    _save_orient(stored, path, orientation)
    return visual.size


class ExifOrientationTests(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.dir = self._td.name

    def tearDown(self):
        self._td.cleanup()

    def test_raw_open_is_landscape_orient_6(self):
        path = os.path.join(self.dir, 'o6.jpg')
        _phone_portrait_jpeg(path, 6)
        with Image.open(path) as raw:
            self.assertEqual(G.exif_orientation_value(raw), 6)
            self.assertGreater(raw.size[0], raw.size[1])

    def test_open_upright_orient_6(self):
        path = os.path.join(self.dir, 'o6.jpg')
        vw, vh = _phone_portrait_jpeg(path, 6)
        up = G.open_upright(path)
        self.assertEqual(up.size, (vw, vh))
        self.assertGreater(up.size[1], up.size[0])
        self.assertGreater(up.getpixel((up.size[0] // 2, 4))[0], 180)
        self.assertGreater(up.getpixel((3, up.size[1] // 2))[1], 180)

    def test_open_upright_orient_8(self):
        path = os.path.join(self.dir, 'o8.jpg')
        vw, vh = _phone_portrait_jpeg(path, 8)
        up = G.open_upright(path)
        self.assertEqual(up.size, (vw, vh))
        self.assertGreater(up.getpixel((up.size[0] // 2, 4))[0], 180)

    def test_exif_upright_not_applied_twice(self):
        path = os.path.join(self.dir, 'o6.jpg')
        _phone_portrait_jpeg(path, 6)
        once = G.open_upright(path)
        twice = G.exif_upright(once)
        self.assertEqual(once.size, twice.size)
        self.assertEqual(once.getpixel((once.size[0] // 2, 4)),
                         twice.getpixel((twice.size[0] // 2, 4)))
        self.assertEqual(G.exif_orientation_value(twice), 1)

    def test_fill_crop_does_not_rotate_90(self):
        path = os.path.join(self.dir, 'o6.jpg')
        _phone_portrait_jpeg(path, 6)
        up = G.open_upright(path).convert('RGB')
        cropped = G.center_crop_to_ar(up, 4 / 3)
        self.assertGreater(cropped.size[0], cropped.size[1])
        left = cropped.getpixel((3, cropped.size[1] // 2))
        top = cropped.getpixel((cropped.size[0] // 2, 3))
        self.assertGreater(left[1], 150)
        self.assertLess(top[1], 120)

    def test_process_photo_fill_stays_upright(self):
        path = os.path.join(self.dir, 'o6.jpg')
        _phone_portrait_jpeg(path, 6)
        out = FX.process_photo(path, src=None, fx={}, box=(400, 300), fit='fill')
        self.assertIsNotNone(out)
        # Ảnh dọc: contain trong ô 4:3, không xoay 90°, không kéo thành 400×300.
        self.assertGreater(out.size[1], out.size[0])
        self.assertLessEqual(out.size[0], 400)
        self.assertLessEqual(out.size[1], 300)
        self.assertNotEqual(out.size, (400, 300))
        left = out.getpixel((3, out.size[1] // 2))
        top = out.getpixel((out.size[0] // 2, 4))
        self.assertGreater(left[1], 120)
        self.assertGreater(top[0], 180)

    def test_process_photo_src_not_transposed_again(self):
        path = os.path.join(self.dir, 'o6.jpg')
        _phone_portrait_jpeg(path, 6)
        src = G.open_upright(path).convert('RGB')
        try:
            src.info.pop('exif', None)
        except Exception:
            pass
        out = FX.process_photo(path, src=src, fx={}, box=(400, 300), fit='fill')
        self.assertGreater(out.size[1], out.size[0])
        self.assertGreater(out.getpixel((3, out.size[1] // 2))[1], 120)
        self.assertGreater(out.getpixel((out.size[0] // 2, 4))[0], 180)

    def test_portrait_fill_contains_full_frame(self):
        path = os.path.join(self.dir, 'portrait.jpg')
        _marked_portrait().save(path, 'JPEG', quality=95)
        out = FX.process_photo(path, src=None, fx={}, box=(400, 300), fit='fill')
        self.assertIsNotNone(out)
        self.assertAlmostEqual(out.size[0] / out.size[1], 0.5, delta=0.05)
        self.assertEqual(out.size[1], 300)
        self.assertGreater(out.getpixel((out.size[0] // 2, 4))[0], 180)
        self.assertGreater(out.getpixel((3, out.size[1] // 2))[1], 150)

    def test_landscape_fill_still_crops_to_box(self):
        path = os.path.join(self.dir, 'land.jpg')
        im = Image.new('RGB', (200, 100), (20, 40, 200))
        for y in range(100):
            for x in range(20):
                im.putpixel((x, y), (20, 200, 40))
        im.save(path, 'JPEG', quality=95)
        out = FX.process_photo(path, src=None, fx={}, box=(400, 300), fit='fill')
        self.assertEqual(out.size, (400, 300))

    def test_photo_fit_mode_portrait_vs_landscape(self):
        self.assertEqual(G.photo_fit_mode('fill', 60, 120), 'fit')
        self.assertEqual(G.photo_fit_mode('fill', 120, 90), 'fill')
        self.assertEqual(G.photo_fit_mode('fit', 120, 90), 'fit')
        self.assertFalse(G.is_portrait_size(100, 100))

    def test_contain_dims_portrait_in_43_box(self):
        w, h = G.contain_dims(60, 120, 4.0, 3.0)
        self.assertAlmostEqual(h, 3.0, places=4)
        self.assertAlmostEqual(w, 1.5, places=4)
        self.assertLess(w, 4.0)

    def test_pptx_photo_path_strips_orientation(self):
        path = os.path.join(self.dir, 'o6.jpg')
        vw, vh = _phone_portrait_jpeg(path, 6)
        baked = EX._pptx_photo_path(path, self.dir)
        self.assertNotEqual(baked, path)
        self.assertEqual(G.exif_orientation(baked), 1)
        with Image.open(baked) as im:
            self.assertEqual(im.size, (vw, vh))

    def test_grid_cells_3_and_4_same_size_no_rotate(self):
        cells = G.grid_cells((0, 0, 8.28, 5.55), 4, 4 / 3, 0.1)
        self.assertEqual(len(cells), 4)
        w0, h0 = cells[0][2], cells[0][3]
        for i in (1, 2, 3):
            self.assertAlmostEqual(cells[i][2], w0)
            self.assertAlmostEqual(cells[i][3], h0)
        self.assertGreater(cells[2][1], cells[0][1])
        self.assertGreater(cells[3][1], cells[1][1])
        # Ô 3–4 cùng tỉ lệ ô 1–2 (gap trừ đều, không xoay).

    def test_pptx_portrait_contained_not_cropped(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        path = os.path.join(self.dir, 'p.jpg')
        _marked_portrait(60, 120).save(path, 'JPEG', quality=95)
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        area = {'x': 0, 'y': 0, 'w': 8.0, 'h': 6.0, 'fit_mode': 'fill',
                'gap': 0.0, 'radius': 0, 'opacity': 100}
        EX._place_images(slide, [path], area, 4 / 3, self.dir, fx={})
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pics), 1)
        pic = pics[0]
        self.assertAlmostEqual(pic.crop_left, 0.0, places=4)
        self.assertAlmostEqual(pic.crop_right, 0.0, places=4)
        self.assertAlmostEqual(pic.crop_top, 0.0, places=4)
        self.assertAlmostEqual(pic.crop_bottom, 0.0, places=4)
        # Ô 4:3 = 8×6; ảnh 1:2 → contain cao 6″, rộng 3″, căn giữa.
        self.assertAlmostEqual(pic.width / 914400, 3.0, places=2)
        self.assertAlmostEqual(pic.height / 914400, 6.0, places=2)
        self.assertGreater(pic.left, 0)


if __name__ == '__main__':
    unittest.main()
