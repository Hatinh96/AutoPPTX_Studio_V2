import os
import tempfile
import unittest
from collections import OrderedDict

from openpyxl import Workbook

from autopptx2.datasource import (
    ExcelSource,
    build_store_index,
    drop_off_groups,
    filter_groups_by_geo,
    inspect_excel,
    looks_like_order_list,
    match_row,
    order_groups,
    sort_codes_by_list_order,
)
from autopptx2.exporter import (
    estimate,
    store_index_page_count,
    write_store_index_xlsx,
)


class TestListOrder(unittest.TestCase):
    def test_filename_and_sheet_detect_list(self):
        self.assertTrue(looks_like_order_list(r'C:\x\List CF.xlsx'))
        self.assertTrue(looks_like_order_list('University List.xlsx'))
        self.assertFalse(looks_like_order_list(
            'master.xlsx', ['Coffee Shop & Milk Tea List']))
        self.assertFalse(looks_like_order_list(
            'W01_Update 06.01.2026_ Golden 2026_ Coffee Shop.xlsx',
            ['Coffee Shop & Milk Tea List']))
        self.assertFalse(looks_like_order_list('Standard_Master_All_Channels.xlsx'))
        self.assertFalse(looks_like_order_list('master.xlsx', ['Data', 'Summary']))

    def test_sort_follows_excel_row_order(self):
        rows = [
            {'Code_RP': 'HLCALMETTE'},
            {'Code_RP': 'HLLIBERTY'},
            {'Code_RP': 'HL119HAMNGHI'},
        ]
        by_code = {r['Code_RP']: r for r in rows}
        self.assertEqual(
            sort_codes_by_list_order(
                ['HL119HAMNGHI', 'HLCALMETTE', 'HLLIBERTY'], rows, by_code),
            ['HLCALMETTE', 'HLLIBERTY', 'HL119HAMNGHI'])

    def test_order_groups_list_mode(self):
        groups = OrderedDict([
            ('HL119HAMNGHI', ['c.jpg']),
            ('HLCALMETTE', ['a.jpg']),
        ])
        rows = [{'Code_RP': 'HLCALMETTE'}, {'Code_RP': 'HL119HAMNGHI'}]
        by_code = {r['Code_RP']: r for r in rows}
        out = order_groups(groups, by_code, excel_rows=rows, sort_mode='list')
        self.assertEqual(list(out), ['HLCALMETTE', 'HL119HAMNGHI'])

    def test_uploaded_list_overrides_master_order(self):
        groups = OrderedDict([
            ('HLLIBERTY', ['b.jpg']),
            ('HLCALMETTE', ['a.jpg']),
            ('HL119HAMNGHI', ['c.jpg']),
        ])
        master = [
            {'Code_RP': 'HL119HAMNGHI'},
            {'Code_RP': 'HLCALMETTE'},
            {'Code_RP': 'HLLIBERTY'},
        ]
        uploaded = [
            {'Code_RP': 'HLCALMETTE'},
            {'Code_RP': 'HLLIBERTY'},
            {'Code_RP': 'HL119HAMNGHI'},
        ]
        by_code = {r['Code_RP']: r for r in master}
        out = order_groups(
            groups, by_code, excel_rows=master, sort_mode='list',
            order_rows=uploaded)
        self.assertEqual(list(out), ['HLCALMETTE', 'HLLIBERTY', 'HL119HAMNGHI'])

    def _write_list_book(self):
        wb = Workbook()
        summary = wb.active
        summary.title = 'Summary'
        summary.append(['Code_RP', 'Name'])
        summary.append(['SHOULD_SKIP', 'Summary row'])
        coffee = wb.create_sheet('Coffee Shop & Milk Tea List')
        coffee.append(['No.', 'Report Code', 'Name'])
        coffee.append(['1', 'HLCALMETTE', 'Highlands Calmette'])
        coffee.append(['2', 'HLLIBERTY', 'Highlands Liberty'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(path)
        wb.close()
        return path

    def test_load_uses_list_sheet_not_summary(self):
        path = self._write_list_book()
        try:
            src = ExcelSource()
            self.assertTrue(src.load(path), src.error)
            self.assertEqual(
                [r['Code_RP'] for r in src.rows],
                ['HLCALMETTE', 'HLLIBERTY'])
            self.assertFalse(looks_like_order_list(path, src.source_sheets))
            info = inspect_excel(path)
            self.assertFalse(info.get('prefer_list_order'))
            self.assertEqual(info.get('n'), 2)
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_off_sheet_excluded_from_active_list(self):
        wb = Workbook()
        listed = wb.active
        listed.title = 'Coffee Shop & Milk Tea List'
        listed.append(['Report Code', 'Name'])
        listed.append(['HLCALMETTE', 'Highlands Calmette'])
        listed.append(['HLVINCOMB3', 'Highlands Vincom B3'])
        off = wb.create_sheet('CỬA HÀNG OFF TRÊN HỆ THỐNG')
        off.append(['Report Code', 'Name'])
        off.append(['HLVINCOMB3', 'Highlands Vincom B3'])
        off.append(['HLBAUCAT', 'Highlands Bàu Cát'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(path)
            src = ExcelSource()
            self.assertTrue(src.load(path), src.error)
            self.assertEqual([r['Code_RP'] for r in src.rows], ['HLCALMETTE'])
            self.assertEqual(src.off_codes, {'HLVINCOMB3', 'HLBAUCAT'})
            self.assertNotIn('HLVINCOMB3', src.by_code())
            self.assertIn('HLCALMETTE', src.by_code())
            all_codes = src.by_code(include_off=True)
            self.assertIn('HLVINCOMB3', all_codes)
            self.assertEqual(all_codes['HLVINCOMB3']['_SiteStatus'], 'off')
            self.assertEqual(all_codes['HLVINCOMB3']['Name'], 'Highlands Vincom B3')
            row, mcode, kind = match_row('HLVINCOMB3DP', all_codes)
            self.assertEqual(kind, 'exact')
            self.assertEqual(mcode, 'HLVINCOMB3')
            self.assertEqual(row['Name'], 'Highlands Vincom B3')
            info = inspect_excel(path)
            self.assertEqual(info.get('n_off'), 2)
            photos = OrderedDict([
                ('HLCALMETTE', ['a.jpg']),
                ('HLVINCOMB3', ['b.jpg']),
                ('HLVINCOMB3DP', ['c.jpg']),
            ])
            self.assertEqual(list(drop_off_groups(photos, src.off_codes)), ['HLCALMETTE'])
            from autopptx2.qa import check as qa_check
            report = qa_check(photos, all_codes)
            self.assertEqual({it.code for it in report.ok},
                             {'HLCALMETTE', 'HLVINCOMB3', 'HLVINCOMB3DP'})
            self.assertFalse(report.unmatched)
            self.assertFalse(any(it.code == 'HLBAUCAT' for it in report.missing))
        finally:
            wb.close()
            if os.path.exists(path):
                os.remove(path)

    def test_off_sheet_remaps_header_when_channel_block_changes(self):
        wb = Workbook()
        listed = wb.active
        listed.title = 'Coffee Shop & Milk Tea List'
        listed.append(['Report Code', 'Name'])
        listed.append(['HLCALMETTE', 'Highlands Calmette'])
        listed.append(['PHIANHHAIR', 'Phi Anh Hair'])
        off = wb.create_sheet('CỬA HÀNG OFF TRÊN HỆ THỐNG')
        off.append(['Chanel', 'Report Code', 'Name', 'Address'])
        off.append(['CF', 'HLVINCOMB3', 'Highlands Vincom B3', '70 Le Thanh Ton'])
        off.append(['BD Code', 'Report Code', 'Name', 'Address'])
        off.append(['HUYBD', 'PHIANHHAIR', 'Phi Anh Hair', '183 Nguyen Thi Dinh'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(path)
            src = ExcelSource()
            self.assertTrue(src.load(path), src.error)
            self.assertEqual(src.off_codes, {'HLVINCOMB3', 'PHIANHHAIR'})
            by_off = {r['Code_RP']: r['Name'] for r in src.off_rows}
            self.assertEqual(by_off['PHIANHHAIR'], 'Phi Anh Hair')
            self.assertEqual([r['Code_RP'] for r in src.rows], ['HLCALMETTE'])
        finally:
            wb.close()
            if os.path.exists(path):
                os.remove(path)

    def test_blank_name_and_channel_headers_like_w01(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Coffee Shop & Milk Tea List'
        ws.append(['COFFEE SHOP & MILK TEA LIST'])
        ws.append(['No.', 'BD Code', '', 'Report Code', '', 'Address', 'City'])
        ws.append(['', '', '', '', '', '', ''])
        ws.append(['1', 'NGANBD', 'CF', 'HLCALMETTE', 'Highlands Calmette',
                   '161 Calmette', 'Hà Nội'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(path)
            src = ExcelSource()
            self.assertTrue(src.load(path), src.error)
            self.assertEqual(src.rows[0]['Code_RP'], 'HLCALMETTE')
            self.assertEqual(src.rows[0]['Name'], 'Highlands Calmette')
            self.assertEqual(src.rows[0]['Channel'], 'CF')
            self.assertEqual(src.rows[0]['City'], 'Hà Nội')
        finally:
            wb.close()
            if os.path.exists(path):
                os.remove(path)

    def test_salon_blank_name_uses_sheet_channel(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Beauty Salon List'
        ws.append(['BEAUTY SALON LIST'])
        ws.append(['No.', 'BD Code', 'Report Code', '', 'Address', 'City'])
        ws.append(['1', 'HABD', 'DUCHAIR', 'Duc Hair', '42 Thu Khoa Huan', 'Hồ Chí Minh'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(path)
            src = ExcelSource()
            self.assertTrue(src.load(path), src.error)
            self.assertEqual(src.rows[0]['Name'], 'Duc Hair')
            self.assertEqual(src.rows[0]['Channel'], 'Beauty Salon')
            self.assertNotEqual(src.rows[0]['Channel'], 'HABD')
        finally:
            wb.close()
            if os.path.exists(path):
                os.remove(path)

    def test_building_premium_channel_from_filename(self):
        wb = Workbook()
        mau = wb.active
        mau.title = 'Mẫu list bán hàng mới'
        mau.append(['DIGITAL BUILDING LIST'])
        mau.append(['No.', 'Report Code', 'Tower', 'Address Detail', 'City'])
        mau.append(['1', 'AKARICITY', 'Akari City', 'Võ Văn Kiệt', 'Hồ Chí Minh'])
        listed = wb.create_sheet('List sale_DIGITAL BUILDING')
        listed.append(['DIGITAL BUILDING LIST'])
        listed.append(['No.', 'Type', 'BD Code', 'OPS', 'CS', 'KT',
                       'Report Code', 'Name of Block', 'Address Detail',
                       'Ward', 'City'])
        listed.append(['BUILDING IN HO CHI MINH'])
        listed.append(['1', 'AP', 'THANHVY', 'HUY', 'NGUYEN', 'TUAN',
                       'AKARICITYBLOCKAK1', 'Akari City Block AK1',
                       'Võ Văn Kiệt', 'Bình Trị Đông', 'Hồ Chí Minh'])
        copy = wb.create_sheet('List sale_DIGITAL BUILDING (1)')
        copy.append(['No.', 'Report Code', 'Name of Block', 'City'])
        copy.append(['1', 'SHOULD_SKIP', 'Copy row', 'Hà Nội'])
        fd, path = tempfile.mkstemp(prefix='List ban hang Building PREMIUM_',
                                    suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(path)
            src = ExcelSource()
            self.assertTrue(src.load(path), src.error)
            self.assertEqual([r['Code_RP'] for r in src.rows],
                             ['AKARICITYBLOCKAK1'])
            self.assertEqual(src.rows[0]['Channel'], 'Building Premium')
            self.assertEqual(src.rows[0]['Name'], 'Akari City Block AK1')
            self.assertEqual(src.source_sheets, ['List sale_DIGITAL BUILDING'])
        finally:
            wb.close()
            if os.path.exists(path):
                os.remove(path)

    def test_store_index_follows_photo_order_not_excel(self):
        groups = OrderedDict([
            ('HL119HAMNGHI', ['c.jpg', 'c2.jpg']),
            ('HLCALMETTE', ['a.jpg']),
        ])
        rows = [
            {'Code_RP': 'HLCALMETTE', 'Name': 'Liberty',
             'District': 'Q1', 'City': 'HCM', 'Channel': 'Hotel'},
            {'Code_RP': 'HL119HAMNGHI', 'Name': '119 Ham Nghi',
             'District': 'Q1', 'City': 'HCM', 'Channel': 'Hotel'},
        ]
        by_code = {r['Code_RP']: r for r in rows}
        idx = build_store_index(groups, by_code)
        self.assertEqual([r['code'] for r in idx],
                         ['HL119HAMNGHI', 'HLCALMETTE'])
        self.assertEqual([r['stt'] for r in idx], [1, 2])
        self.assertEqual(idx[0]['photos'], 2)
        self.assertEqual(idx[0]['name'], '119 Ham Nghi')
        self.assertEqual(idx[1]['name'], 'Liberty')
        self.assertIn(idx[0]['status'], ('thieu', 'du', 'thua', 'chua_excel', 'chua_man'))

    def test_store_index_xlsx_and_estimate_pages(self):
        self.assertEqual(store_index_page_count(0), 0)
        self.assertEqual(store_index_page_count(18), 1)
        self.assertEqual(store_index_page_count(19), 2)
        groups = {f'C{i}': ['a.jpg'] for i in range(18)}
        _ng, ns_off, _ = estimate(groups, 1, store_index=False)
        _ng, ns_on, _ = estimate(groups, 1, store_index=True)
        self.assertEqual(ns_off, 18)
        self.assertEqual(ns_on, 19)
        rows = build_store_index(
            OrderedDict([('HL1', ['a.jpg']), ('HL2', ['b.jpg'])]), {})
        fd, path = tempfile.mkstemp(prefix='idx_', suffix='.xlsx')
        os.close(fd)
        try:
            write_store_index_xlsx(path, rows)
            from openpyxl import load_workbook
            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws['A1'].value, 'STT')
            self.assertEqual(ws['A2'].value, 1)
            self.assertEqual(ws['B2'].value, 'HL1')
            self.assertEqual(ws['A3'].value, 2)
            wb.close()
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_fuzzy_does_not_attach_live_code_to_off(self):
        by_code = {
            'HLVINCOM': {'_SiteStatus': 'on', 'Name': 'Vincom sống',
                         'Channel': 'Coffee Shop & Milk Tea'},
            'HLVINCOMB3': {'_SiteStatus': 'off', 'Name': 'Vincom off',
                           'Channel': 'Coffee Shop & Milk Tea'},
        }
        row, code, kind = match_row('HLVINCOMB', by_code)
        self.assertNotEqual(code, 'HLVINCOMB3')
        if kind == 'fuzzy':
            self.assertEqual(str(row.get('_SiteStatus')), 'on')

    def test_channel_filter_drops_other_channel_photos(self):
        groups = OrderedDict([
            ('UNI1', ['u.jpg']),
            ('CF1', ['c.jpg']),
        ])
        by_code = {
            'UNI1': {'Channel': 'University', 'City': 'Hà Nội', 'District': ''},
            'CF1': {'Channel': 'Coffee Shop & Milk Tea', 'City': 'Hà Nội',
                    'District': ''},
        }
        out = filter_groups_by_geo(
            groups, by_code, channel='University',
            all_channels='Tất cả kênh')
        self.assertEqual(list(out), ['UNI1'])


if __name__ == '__main__':
    unittest.main()
