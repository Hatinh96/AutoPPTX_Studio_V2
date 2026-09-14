"""Phân tích THE NELSON.pptx → layout 16:9 cho AutoPPTX."""
from __future__ import annotations

import json
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SRC_W, SRC_H = 20.0, 11.25
DST_W, DST_H = 13.33, 7.5
SX, SY = DST_W / SRC_W, DST_H / SRC_H


def emu_in(v):
    return round(v / 914400, 4) if v is not None else None


def scaled(v, axis):
    if v is None:
        return None
    return round(v * (SX if axis == "x" else SY), 3)


def rgb(c):
    try:
        if c and c.type is not None:
            return str(c.rgb)
    except Exception:
        pass
    return None


def shape_info(shape, depth=0):
    l, t, w, h = (
        emu_in(shape.left),
        emu_in(shape.top),
        emu_in(shape.width),
        emu_in(shape.height),
    )
    info = {
        "name": shape.name,
        "type": int(shape.shape_type),
        "depth": depth,
        "in": {"x": l, "y": t, "w": w, "h": h},
        "layout_16x9": {
            "x": scaled(l, "x"),
            "y": scaled(t, "y"),
            "w": scaled(w, "x"),
            "h": scaled(h, "y"),
        },
    }
    if shape.has_text_frame:
        info["text"] = shape.text_frame.text.replace("\n", " | ")[:200]
        for p in shape.text_frame.paragraphs[:1]:
            if p.runs:
                r = p.runs[0]
                info["font"] = {
                    "name": r.font.name,
                    "size": r.font.size.pt if r.font.size else None,
                    "bold": r.font.bold,
                    "color": rgb(r.font.color),
                }
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        info["picture"] = True
    if shape.has_table:
        tbl = shape.table
        info["table"] = f"{len(tbl.rows)}x{len(tbl.columns)}"
        cells = []
        for ri in range(min(5, len(tbl.rows))):
            row = [tbl.cell(ri, ci).text[:40] for ci in range(len(tbl.columns))]
            cells.append(row)
        info["cells"] = cells
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        info["children"] = []
        for ch in shape.shapes:
            info["children"].append(shape_info(ch, depth + 1))
    return info


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else r"THE NELSON.pptx"
    prs = Presentation(path)
    sw = prs.slide_width / 914400
    sh = prs.slide_height / 914400
    out = {
        "file": path,
        "size_in": [sw, sh],
        "scale_to_16x9": [SX, SY],
        "slides": [],
    }
    for si in range(len(prs.slides)):
        slide = prs.slides[si]
        shapes = [shape_info(s) for s in slide.shapes]
        out["slides"].append({"index": si + 1, "shape_count": len(shapes), "shapes": shapes})
    text = json.dumps(out, ensure_ascii=False, indent=2)
    out_path = sys.argv[2] if len(sys.argv) > 2 else None
    if out_path:
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(text)
    else:
        sys.stdout.buffer.write(text.encode("utf-8"))


if __name__ == "__main__":
    main()
