"""Đối chiếu list sales × ảnh lúc xuất PPTX.

Quy tắc ĐỦ: mã trên list (phạm vi kênh/quận của ảnh đã khớp) có ≥ 1 ảnh.
Không so với số tivi. Không import Tkinter.
"""
from dataclasses import dataclass, field

from .datasource import match_row, clean


def _norm(v):
    return str(clean(v) or '').strip().casefold()


@dataclass
class QAItem:
    code: str
    excel_code: str = ''
    name: str = ''
    district: str = ''
    channel: str = ''
    photos: int = 0
    status: str = ''      # ok | missing | unmatched | fuzzy
    note: str = ''


@dataclass
class QAReport:
    missing: list = field(default_factory=list)      # list có, chưa có ảnh
    unmatched: list = field(default_factory=list)    # ảnh không khớp mã Excel
    fuzzy: list = field(default_factory=list)        # ảnh khớp gần đúng
    ok: list = field(default_factory=list)
    list_count: int = 0
    photo_groups: int = 0
    scope_note: str = ''
    has_excel: bool = False

    def needs_gate(self):
        return bool(self.missing or self.unmatched or self.fuzzy)


def check(groups, by_code, merged=None):
    """groups: {mã ảnh: [paths]}, by_code: {MÃ EXCEL: row} đã lọc kênh."""
    groups = groups or {}
    by_code = by_code or {}
    rep = QAReport(list_count=len(by_code), photo_groups=len(groups),
                   has_excel=bool(by_code))
    if not by_code:
        rep.scope_note = (
            'Chưa có FILE TỔNG / MASTER đang mở — bỏ qua đối chiếu. '
            'Đồng bộ hoặc chọn Excel ở tab Nguồn.')
        return rep

    covered = set()
    matched_rows = []
    for code, paths in groups.items():
        n = len(paths or [])
        row, mcode, kind = match_row(code, by_code, merged)
        name = str(clean((row or {}).get('Name'))).strip()
        dist = str(clean((row or {}).get('District'))).strip()
        ch = str(clean((row or {}).get('Channel'))).strip()
        item = QAItem(
            code=str(code), excel_code=str(mcode or ''),
            name=name or str(code), district=dist, channel=ch, photos=n)

        if kind == 'none':
            item.status = 'unmatched'
            item.note = 'Tên file không khớp mã Excel — dễ thiếu ảo.'
            item.name = str(code)
            rep.unmatched.append(item)
            continue

        extra = (row or {}).get('_merged_codes') or []
        for c in extra:
            covered.add(str(c).strip().upper())
        if mcode:
            covered.add(str(mcode).strip().upper())
        covered.add(str(code).strip().upper())
        matched_rows.append(row or {})

        if kind == 'fuzzy':
            item.status = 'fuzzy'
            item.note = f'Gần đúng với Excel {mcode} — kiểm tra tên file.'
            rep.fuzzy.append(item)
        else:
            item.status = 'ok'
            item.note = 'Đủ ảnh'
            rep.ok.append(item)

    scope_ch = {_norm(r.get('Channel')) for r in matched_rows if _norm(r.get('Channel'))}
    scope_dt = {_norm(r.get('District')) for r in matched_rows if _norm(r.get('District'))}
    bits = []
    if scope_ch:
        bits.append('kênh ' + ', '.join(sorted(scope_ch)))
    if scope_dt:
        bits.append('quận/TP ' + ', '.join(sorted(scope_dt)))
    rep.scope_note = ('Phạm vi suy từ ảnh đã khớp: ' + ' · '.join(bits)
                      if bits else 'Phạm vi: toàn bộ list đang lọc.')

    for k, row in by_code.items():
        ku = str(k).strip().upper()
        if ku in covered:
            continue
        if scope_ch and _norm(row.get('Channel')) not in scope_ch:
            continue
        if scope_dt and _norm(row.get('District')) not in scope_dt:
            continue
        rep.missing.append(QAItem(
            code=ku, excel_code=ku,
            name=str(clean(row.get('Name'))).strip() or ku,
            district=str(clean(row.get('District'))).strip(),
            channel=str(clean(row.get('Channel'))).strip(),
            photos=0, status='missing',
            note='Có trên list, chưa có ảnh.'))
    rep.missing.sort(key=lambda x: x.code)
    return rep


def checklist_path(pptx_path):
    root, _ = os_path_splitext(pptx_path)
    return root + '_checklist.xlsx'


def os_path_splitext(p):
    import os
    return os.path.splitext(p or '')


def write_checklist(path, report: QAReport):
    """Ghi file Excel nội bộ — khách hàng không thấy trong PPTX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Đối chiếu'

    headers = ('Kết luận', 'Mã ảnh', 'Mã Excel', 'Tên địa điểm', 'Quận / TP',
               'Kênh', 'Số ảnh', 'Ghi chú')
    fills = {
        'missing': PatternFill('solid', fgColor='FECACA'),
        'unmatched': PatternFill('solid', fgColor='FDE68A'),
        'fuzzy': PatternFill('solid', fgColor='FDE68A'),
        'ok': PatternFill('solid', fgColor='BBF7D0'),
    }
    labels = {
        'missing': 'THIẾU ẢNH',
        'unmatched': 'LỆCH MÃ FILE',
        'fuzzy': 'GẦN ĐÚNG',
        'ok': 'ĐỦ',
    }
    thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'))
    head_fill = PatternFill('solid', fgColor='111827')
    head_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
    body_font = Font(name='Calibri', size=11)

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(1, col)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal='center', wrap_text=True)

    rows = list(report.missing) + list(report.unmatched) + list(report.fuzzy) + list(report.ok)
    for i, it in enumerate(rows, 2):
        vals = (
            labels.get(it.status, it.status),
            it.code, it.excel_code, it.name, it.district,
            it.channel, it.photos, it.note,
        )
        for col, v in enumerate(vals, 1):
            cell = ws.cell(i, col, v)
            cell.font = body_font
            cell.border = thin
            fill = fills.get(it.status)
            if fill:
                cell.fill = fill
            if col == 7:
                cell.alignment = Alignment(horizontal='center')

    widths = (16, 18, 18, 36, 22, 22, 10, 48)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f'A1:H{max(1, len(rows) + 1)}'
    ws.freeze_panes = 'A2'

    sm = wb.create_sheet('Tổng', 0)
    sm.append(('Hạng mục', 'Số'))
    sm['A1'].font = head_font
    sm['B1'].font = head_font
    sm['A1'].fill = head_fill
    sm['B1'].fill = head_fill
    summary = [
        ('List đang lọc (Excel)', report.list_count),
        ('Nhóm ảnh', report.photo_groups),
        ('Đủ ảnh', len(report.ok)),
        ('List chưa có ảnh (phạm vi báo cáo)', len(report.missing)),
        ('Ảnh không khớp mã Excel', len(report.unmatched)),
        ('Ảnh khớp gần đúng', len(report.fuzzy)),
        ('Phạm vi', report.scope_note),
    ]
    for row in summary:
        sm.append(row)
    sm.column_dimensions['A'].width = 42
    sm.column_dimensions['B'].width = 70

    wb.save(path)
    return path
