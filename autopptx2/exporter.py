"""Sinh file PPTX — module THUẦN dữ liệu, không import Tkinter.

Chạy trong worker thread; giao tiếp với UI qua callback (progress, log)
và threading.Event (huỷ). Mọi thông số được chụp sẵn vào ExportJob trên
main thread trước khi chạy.
"""
import os
import re
import math
import time
import shutil
import hashlib
import tempfile
from dataclasses import dataclass, field
import datetime

from PIL import Image, ImageDraw
from pptx.enum.shapes import MSO_SHAPE
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

from .constants import (SLIDE_W_IN, SLIDE_H_IN,
                        TABLE_HEADER_BG, TABLE_HEADER_FG, TABLE_DATA_BG,
                        TABLE_SUB_FG, TABLE_LINE, TABLE_TEXT,
                        NELSON_HEADER_BG, NELSON_HEADER_FG, NELSON_DATA_BG,
                        NELSON_LINE, NELSON_TEXT)
from . import geometry as G
from . import effects as FX
from . import fonts as FN
from .datasource import (match_row, build_merged_groups, build_info_rows,
                         build_info_table, build_nelson_table, channel_text,
                         clean, build_overview, wrap_info_cell, info_row_weights,
                         _info_value_wrap_chars, _address_line,
                         location_fx_fields,
                         order_groups, filter_groups_by_geo,
                         group_export_chunks, screen_qty)


class ExportCancelled(Exception):
    pass


@dataclass
class ExportJob:
    layout: dict                      # snapshot bố cục (deep copy)
    groups: dict                      # OrderedDict {mã: [đường dẫn ảnh tuyệt đối]}
    excel_by_code: dict               # {MÃ: row} (đã lọc kênh)
    avatar_map: dict                  # {MÃ: đường dẫn avatar} dựng sẵn ở main thread
    out_path: str
    n_per_slide: int = 4
    slides_per_file: int = 0          # 0 = một file; >0 = cắt mỗi N slide
    img_ar: float | None = 4 / 3      # None = Tự do
    bg_image: str | None = None
    channel_enabled: bool = True
    channel_template: str = "Report {channel} 2026"
    visible: dict = field(default_factory=dict)   # {tên phần tử: bool}
    fx: dict = field(default_factory=dict)        # đóng dấu / mini-map / AI
    excel_rows: list = field(default_factory=list)  # mọi dòng (khu/loại màn)
    slide_style: str = 'report'                   # report | saleskit | nelson
    pad_blank_slides: bool = False                # slide trắng / ô trống theo số màn
    sort_mode: str = 'city'                       # city | list
    sort_rows: list = field(default_factory=list)  # list up riêng để xếp STT
    city_filter: str = ''
    district_filter: str = ''
    split_export_by: str = 'none'                 # none | city | district
    export_pdf: bool = False


@dataclass
class ExportReport:
    files: list = field(default_factory=list)
    slides: int = 0
    groups: int = 0
    fuzzy: list = field(default_factory=list)       # (mã ảnh, mã Excel gần đúng)
    unmatched: list = field(default_factory=list)   # mã ảnh không có dòng Excel
    missing_avatars: list = field(default_factory=list)
    seconds: float = 0.0
    cancelled: bool = False


def estimate(groups, n_per_slide, overview_rows=0, slides_per_file=0,
             by_code=None, merged=None, pad_blank=False):
    """(số nhóm, số slide dự kiến, số file dự kiến)."""
    by_code = by_code or {}
    if merged is None and by_code:
        merged = build_merged_groups(by_code)
    slides = 0
    for code, paths in (groups or {}).items():
        row, _, _ = match_row(code, by_code, merged)
        sq = screen_qty(row) if pad_blank else 0
        slides += G.slides_for_group(paths, n_per_slide, sq, pad_blank)
    if overview_rows:
        slides += overview_page_count(overview_rows)
    parts = G.file_part_count(slides, slides_per_file)
    return len(groups or {}), slides, parts


def _safe_filename_part(text, max_len=48):
    s = re.sub(r'[<>:"/\\|?*]+', '_', str(text or '').strip())
    s = re.sub(r'\s+', '_', s).strip('._')
    return (s[:max_len] if s else 'Export')


OV_FIRST = 12
OV_NEXT = 16


def overview_page_count(n_rows):
    n_rows = max(0, int(n_rows or 0))
    if n_rows <= OV_FIRST:
        return 1
    return 1 + math.ceil((n_rows - OV_FIRST) / OV_NEXT)


# ────────────────────────────────────────────────────────────
def _new_prs():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def _blank_layout(prs):
    for lo in prs.slide_layouts:
        if lo.name.lower() in ('blank', 'trống'):
            return lo
    return prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]


def _hex_rgb(hexv, fallback=(0, 0, 0)):
    try:
        h = str(hexv).lstrip('#')
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except Exception:
        return RGBColor(*fallback)


def _p_run(p):
    """Run đầu tiên của đoạn; tạo run rỗng nếu chưa có (đoạn text='')."""
    return p.runs[0] if p.runs else p.add_run()


def _add_bg(slide, prs, bg_path):
    pic = slide.shapes.add_picture(bg_path, 0, 0,
                                   width=prs.slide_width, height=prs.slide_height)
    # đưa xuống đáy — mọi phần tử khác nằm trên
    sp = pic._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def _style_fit_textbox(tb, text, width_in, height_in, default_size,
                       font_name='Arial', bold=True, italic=False, rgb=None,
                       align='left', max_lines=2, min_size=9, anchor='middle',
                       opacity=100, tracking=0):
    """Đổ text + tự co cỡ chữ (dùng chung công thức với preview)."""
    tf = tb.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    try:
        tf.margin_left = Pt(2)
        tf.margin_right = Pt(2)
        tf.margin_top = Pt(0)
        tf.margin_bottom = Pt(0)
        tf.vertical_anchor = {'top': MSO_ANCHOR.TOP,
                              'middle': MSO_ANCHOR.MIDDLE,
                              'bottom': MSO_ANCHOR.BOTTOM}.get(anchor,
                                                               MSO_ANCHOR.MIDDLE)
    except Exception:
        pass
    size = G.fit_block_size(text, width_in, height_in, default_size,
                            max_lines=max_lines, min_size=min_size)
    lines = str(text or "").split("\n")
    tf.text = lines[0]
    for extra in lines[1:]:
        tf.add_paragraph().text = extra
    amap = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
            'right': PP_ALIGN.RIGHT}
    for para in tf.paragraphs:
        para.alignment = amap.get(align, PP_ALIGN.LEFT)
        para.font.name = FN.safe_family(font_name, text)
        para.font.bold = bold
        para.font.italic = italic
        para.font.size = Pt(size)
        if rgb is not None:
            para.font.color.rgb = rgb
        for run in para.runs:
            _apply_run_alpha(run, opacity)
            _apply_run_tracking(run, tracking)
    return size


def _apply_run_alpha(run, opacity):
    """Độ mờ chữ 0–100 → DrawingML a:alpha (100000 = đặc)."""
    try:
        op = float(opacity if opacity is not None else 100)
    except Exception:
        op = 100
    if op >= 99.5:
        return
    try:
        from lxml import etree
        from pptx.oxml.ns import qn
        rPr = run._r.get_or_add_rPr()
        solid = rPr.find(qn('a:solidFill'))
        if solid is None:
            return
        srgb = solid.find(qn('a:srgbClr'))
        if srgb is None:
            return
        for old in srgb.findall(qn('a:alpha')):
            srgb.remove(old)
        el = etree.SubElement(srgb, qn('a:alpha'))
        el.set('val', str(max(0, min(100000, int(round(op * 1000))))))
    except Exception:
        pass


def _apply_run_tracking(run, spc_pt):
    try:
        spc = float(spc_pt or 0)
    except Exception:
        return
    if abs(spc) < 0.01:
        return
    try:
        rPr = run._r.get_or_add_rPr()
        rPr.set('spc', str(int(round(spc * 100))))
    except Exception:
        pass


def _set_pic_opacity(pic, opacity):
    try:
        op = float(opacity if opacity is not None else 100)
    except Exception:
        op = 100
    if op >= 99.5:
        return
    try:
        from lxml import etree
        from pptx.oxml.ns import qn
        blip = pic._element.find('.//' + qn('a:blip'))
        if blip is None:
            return
        for old in blip.findall(qn('a:alphaModFix')):
            blip.remove(old)
        el = etree.SubElement(blip, qn('a:alphaModFix'))
        el.set('amt', str(max(0, min(100000, int(round(op * 1000))))))
    except Exception:
        pass


def _add_info_table(slide, info_rows, ninfo, font_cfg):
    """Khối Thông tin mẫu báo cáo = BẢNG PowerPoint 2 cột."""
    from pptx.oxml.ns import qn
    rows = [(l, str(v), acc) for (l, v, acc) in info_rows]
    if not rows:
        return
    x, y, w, h = ninfo['x'], ninfo['y'], ninfo['w'], ninfo['h']
    gt = slide.shapes.add_table(len(rows), 2, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    table = gt.table
    table.first_row = False
    table.horz_banding = False
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tblPr.append(sid)
    table.columns[0].width = Inches(w * 0.42)
    table.columns[1].width = Inches(w * 0.58)

    fname = ninfo.get('font') or font_cfg.get('name', 'Arial')
    vr = _hex_rgb(font_cfg.get('color') or '#0f172a', (15, 23, 42))
    label_rgb = RGBColor(*[int(c + (255 - c) * 0.45) for c in (vr[0], vr[1], vr[2])])
    accent_rgb = _hex_rgb(ninfo.get('accent_color') or '#2563eb', (37, 99, 235))
    size_pt = float(ninfo.get('size', 12))
    op = ninfo.get('opacity', 100)

    weights = info_row_weights(rows, w)
    tot_w = sum(weights) or 1.0
    wrap_chars = _info_value_wrap_chars(w)

    for i, wt in enumerate(weights):
        try:
            table.rows[i].height = Inches(max(0.16, h * wt / tot_w))
        except Exception:
            pass
    for i, (label, value, accent) in enumerate(rows):
        c0, c1 = table.cell(i, 0), table.cell(i, 1)
        for c in (c0, c1):
            c.fill.background()
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.02)
            c.margin_right = Inches(0.02)
            c.margin_top = Inches(0.0)
            c.margin_bottom = Inches(0.0)
        c1.text_frame.word_wrap = True
        p0 = c0.text_frame.paragraphs[0]
        p0.text = label
        p0.alignment = PP_ALIGN.LEFT
        f0 = _p_run(p0).font
        f0.name = FN.safe_family(fname, label)
        f0.size = Pt(max(6, size_pt * 0.85))
        f0.bold = False
        f0.color.rgb = label_rgb
        _apply_run_alpha(_p_run(p0), op)
        tf1 = c1.text_frame
        lines = (wrap_info_cell(value, wrap_chars) if label == 'Address'
                 else [str(value)])
        tf1.clear()
        for j, line in enumerate(lines):
            p1 = tf1.paragraphs[0] if j == 0 else tf1.add_paragraph()
            p1.text = line
            p1.alignment = PP_ALIGN.RIGHT
            f1 = _p_run(p1).font
            f1.name = FN.safe_family(fname, line)
            f1.bold = True
            f1.size = Pt(max(7, size_pt * (1.5 if accent else 1.0)))
            f1.color.rgb = accent_rgb if accent else vr
            _apply_run_alpha(_p_run(p1), op)


def _add_saleskit_table(slide, info_table, ninfo, font_cfg):
    """Khung bảng SALESKIT (5 cột, gộp ô ĐỊA ĐIỂM/ĐỊA CHỈ/TRAFFIC)."""
    from lxml import etree
    from pptx.oxml.ns import qn
    table = info_table or {}
    specs = list(table.get('specs') or [{'area': '', 'form': '', 'size': '',
                                         'qty': '', 'note': ''}])
    n_rows = 3 + max(1, len(specs))
    n_cols = 5
    x, y, w, h = ninfo['x'], ninfo['y'], ninfo['w'], ninfo['h']
    gt = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tbl = gt.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tblPr.append(sid)
    fracs = (0.22, 0.18, 0.22, 0.16, 0.22)
    for i, f in enumerate(fracs):
        tbl.columns[i].width = Inches(w * f)
    row_h = Inches(max(0.16, h / n_rows))
    for i in range(n_rows):
        try:
            tbl.rows[i].height = row_h
        except Exception:
            pass
    try:
        tbl.cell(0, 0).merge(tbl.cell(0, 1))
        tbl.cell(0, 2).merge(tbl.cell(0, 3))
        tbl.cell(1, 0).merge(tbl.cell(1, 1))
        tbl.cell(1, 2).merge(tbl.cell(1, 3))
    except Exception:
        pass

    raw_name = ninfo.get('font') or font_cfg.get('name', 'Arial') or 'Arial'
    size_pt = float(ninfo.get('size', 11))
    hdr_bg = ninfo.get('accent_color') or TABLE_HEADER_BG
    op = ninfo.get('opacity', 100)

    def paint(r, c, text, *, bg, fg, bold, align='left', sz=None):
        cell = tbl.cell(r, c)
        _fill_cell(cell, bg)
        _border_cell(cell, TABLE_LINE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.04)
        cell.margin_right = Inches(0.04)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = str(text or '')
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                       'right': PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        run = _p_run(p)
        run.font.name = FN.safe_family(raw_name, text)
        run.font.bold = bold
        run.font.size = Pt(max(7, sz if sz is not None else size_pt))
        run.font.color.rgb = _hex_rgb(fg)
        _apply_run_alpha(run, op)

    paint(0, 0, 'ĐỊA ĐIỂM', bg=hdr_bg, fg=TABLE_HEADER_FG, bold=True, align='center')
    paint(0, 2, 'ĐỊA CHỈ', bg=hdr_bg, fg=TABLE_HEADER_FG, bold=True, align='center')
    paint(0, 4, 'TRAFFIC', bg=hdr_bg, fg=TABLE_HEADER_FG, bold=True, align='center')
    paint(1, 0, table.get('school') or '', bg=TABLE_DATA_BG, fg=TABLE_TEXT, bold=True)
    paint(1, 2, table.get('address') or '', bg=TABLE_DATA_BG, fg=TABLE_TEXT,
          bold=False, sz=size_pt * 0.9)
    paint(1, 4, table.get('traffic') or '', bg=TABLE_DATA_BG, fg=TABLE_TEXT,
          bold=True, align='center')
    sub_h = ('Khu vực', 'Hình thức', 'Kích thước', 'Số lượng', 'Note')
    for c, lab in enumerate(sub_h):
        paint(2, c, lab, bg='#FFFFFF', fg=hdr_bg, bold=True, align='center',
              sz=size_pt * 0.85)
    keys = ('area', 'form', 'size', 'qty', 'note')
    aligns = ('left', 'center', 'center', 'center', 'left')
    for i, spec in enumerate(specs):
        bg = TABLE_DATA_BG if i % 2 == 0 else '#FFFFFF'
        for c, key in enumerate(keys):
            paint(3 + i, c, spec.get(key) or '', bg=bg, fg=TABLE_TEXT,
                  bold=False, align=aligns[c], sz=size_pt * 0.9)


def _add_nelson_table(slide, nelson_table, ninfo, font_cfg):
    """Bảng NELSON: Số lượng | Note | Hình thức."""
    table = nelson_table or {}
    specs = list(table.get('specs') or [{'qty': '', 'note': '', 'form': ''}])
    n_rows = 1 + max(1, len(specs))
    n_cols = 3
    x, y, w, h = ninfo['x'], ninfo['y'], ninfo['w'], ninfo['h']
    gt = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tbl = gt.table
    tbl.first_row = False
    tbl.horz_banding = False
    from lxml import etree
    from pptx.oxml.ns import qn
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tblPr.append(sid)
    fracs = (0.18, 0.34, 0.48)
    for i, f in enumerate(fracs):
        tbl.columns[i].width = Inches(w * f)
    row_h = Inches(max(0.16, h / n_rows))
    for i in range(n_rows):
        try:
            tbl.rows[i].height = row_h
        except Exception:
            pass

    raw_name = ninfo.get('font') or font_cfg.get('name', 'Arial') or 'Arial'
    size_pt = float(ninfo.get('size', 11))
    hdr_bg = ninfo.get('accent_color') or NELSON_HEADER_BG
    op = ninfo.get('opacity', 100)

    def paint(r, c, text, *, bg, fg, bold, align='left', sz=None):
        cell = tbl.cell(r, c)
        _fill_cell(cell, bg)
        _border_cell(cell, NELSON_LINE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.04)
        cell.margin_right = Inches(0.04)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = str(text or '')
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                       'right': PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        run = _p_run(p)
        run.font.name = FN.safe_family(raw_name, text)
        run.font.bold = bold
        run.font.size = Pt(max(7, sz if sz is not None else size_pt))
        run.font.color.rgb = _hex_rgb(fg)
        _apply_run_alpha(run, op)

    headers = ('Số lượng', 'Note', 'Hình thức')
    aligns = ('center', 'left', 'left')
    for c, lab in enumerate(headers):
        paint(0, c, lab, bg=hdr_bg, fg=NELSON_HEADER_FG, bold=True,
              align=aligns[c], sz=size_pt * 0.9)
    keys = ('qty', 'note', 'form')
    for i, spec in enumerate(specs):
        bg = NELSON_DATA_BG if i % 2 == 0 else '#FFFFFF'
        for c, key in enumerate(keys):
            paint(1 + i, c, spec.get(key) or '', bg=bg, fg=NELSON_TEXT,
                  bold=False, align=aligns[c], sz=size_pt * 0.9)


def _fill_cell(cell, hexv):
    try:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _hex_rgb(hexv, (255, 255, 255))
    except Exception:
        pass


def _border_cell(cell, hexv, width_pt=0.75):
    try:
        from lxml import etree
        from pptx.oxml.ns import qn
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for edge in ('lnL', 'lnR', 'lnT', 'lnB'):
            for old in tcPr.findall(qn('a:' + edge)):
                tcPr.remove(old)
            ln = etree.SubElement(tcPr, qn('a:' + edge))
            ln.set('w', str(int(width_pt * 12700)))
            sf = etree.SubElement(ln, qn('a:solidFill'))
            srgb = etree.SubElement(sf, qn('a:srgbClr'))
            srgb.set('val', str(hexv).lstrip('#').upper())
    except Exception:
        pass


def _pptx_photo_path(path, temp_dir):
    """File chèn PPTX: pixel đã đúng chiều, không EXIF Orientation 2–8.

    PowerPoint tự xoay JPEG theo tag — nếu PIL đã tính crop theo ảnh đứng
    mà file gốc vẫn còn Orientation=6/8 thì slide bị xoay thêm 90°.
    """
    if not path or G.exif_orientation(path) == 1:
        return path
    try:
        key = hashlib.md5(
            f"upright|{os.path.abspath(path)}|{os.path.getmtime(path)}".encode(
                errors='replace')).hexdigest()
    except Exception:
        key = hashlib.md5(f"upright|{path}".encode(errors='replace')).hexdigest()
    out = os.path.join(temp_dir, f"up_{key}.jpg")
    if os.path.exists(out):
        return out
    im = G.open_upright(path)
    try:
        rgb = im.convert('RGB')
        try:
            rgb.info.pop('exif', None)
        except Exception:
            pass
        rgb.save(out, 'JPEG', quality=95)
        return out
    finally:
        try:
            im.close()
        except Exception:
            pass


def _place_avatar(slide, path, av, temp_dir):
    """Avatar center-crop lấp khung `ar` bằng crop_* của PPTX (không méo)."""
    a_ar = av.get('ar') or (4 / 3)
    w_in = av['w']
    h_in = w_in / a_ar
    src = _pptx_photo_path(path, temp_dir)
    pic = slide.shapes.add_picture(src, Inches(av['x']), Inches(av['y']),
                                   width=Inches(w_in), height=Inches(h_in))
    try:
        iw, ih = G.image_size_upright(path)
        l, r, t, b = G.crop_fractions((iw / ih) if ih else a_ar, a_ar)
        pic.crop_left, pic.crop_right = l, r
        pic.crop_top, pic.crop_bottom = t, b
    except Exception:
        pass
    _set_pic_opacity(pic, av.get('opacity', 100))
    return pic


def _rounded_png(src_path, box_ar, radius_pct, temp_dir):
    """Ảnh bo góc: crop-fill về tỉ lệ ô + alpha bo góc → PNG tạm (cache theo
    nội dung tham số). Chỉ dùng khi radius > 0 (đường chậm)."""
    key = hashlib.md5(f"{src_path}|{os.path.getmtime(src_path)}|"
                      f"{box_ar:.4f}|{radius_pct}".encode()).hexdigest()
    out = os.path.join(temp_dir, f"r_{key}.png")
    if os.path.exists(out):
        return out
    im = Image.open(src_path)
    try:
        im.draft('RGB', (2048, 2048))
    except Exception:
        pass
    im = G.exif_upright(im)
    im = im.convert("RGB")
    im = G.center_crop_to_ar(im, box_ar)
    if im.width > 1800:
        im = im.resize((1800, max(1, int(1800 / box_ar))),
                       Image.Resampling.LANCZOS)
    im = im.convert("RGBA")
    rad = int((radius_pct / 100) * (min(im.size) / 2))
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, im.width, im.height],
                                           radius=rad, fill=255)
    im.putalpha(mask)
    im.save(out, "PNG")
    return out


def _add_contained_picture(slide, src, bx, by, bw, bh, iar):
    """Chèn ảnh vừa khít ô (letterbox), không crop, không kéo méo."""
    w, h = G.contain_dims(iar, 1.0, bw, bh)
    return slide.shapes.add_picture(
        src, Inches(bx + (bw - w) / 2), Inches(by + (bh - h) / 2),
        width=Inches(w), height=Inches(h))


def _place_images(slide, paths, area_cfg, img_ar, temp_dir, fx=None):
    """Xếp lưới ảnh trong vùng — None trong paths = ô trống (slide trắng)."""
    paths = list(paths or [])
    num = len(paths)
    if num == 0:
        return
    area = (area_cfg['x'], area_cfg['y'], area_cfg['w'],
            area_cfg.get('h', area_cfg['w'] * 0.66))
    gap = area_cfg.get('gap', 0.1)
    fit_mode = area_cfg.get('fit_mode', 'fill')
    radius = int(area_cfg.get('radius', 0) or 0)
    opacity = area_cfg.get('opacity', 100)
    fx = fx or {}
    bake_fx = FX.needed(fx)
    cells = G.grid_cells(area, num, img_ar, gap)

    for path, (bx, by, bw, bh) in zip(paths, cells):
        if not path:
            try:
                sh = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, Inches(bx), Inches(by),
                    Inches(bw), Inches(bh))
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor(245, 245, 245)
                sh.line.color.rgb = RGBColor(210, 210, 210)
            except Exception:
                pass
            continue
        iw = ih = 0
        try:
            iw, ih = G.image_size_upright(path)
            iar = (iw / ih) if ih else 1.0
        except Exception:
            iar = 1.0
        use_fit = G.photo_fit_mode(fit_mode, iw, ih)
        box_ar = bw / bh if bh else 1.0
        pic = None
        baked = False

        if bake_fx:
            try:
                rp = FX.bake_export_png(path, fx, box_ar, use_fit, radius,
                                        temp_dir, opacity)
                if rp:
                    if use_fit == 'fill':
                        pic = slide.shapes.add_picture(
                            rp, Inches(bx), Inches(by),
                            width=Inches(bw), height=Inches(bh))
                    else:
                        try:
                            with Image.open(rp) as im:
                                piw, pih = im.size
                            par = (piw / pih) if pih else 1.0
                        except Exception:
                            par = iar
                        pic = _add_contained_picture(
                            slide, rp, bx, by, bw, bh, par)
                    baked = True
            except Exception as e:
                print('Bake image err:', e)
                pic = None

        if pic is None and use_fit == 'fill':
            if radius > 0:
                try:
                    rp = _rounded_png(path, box_ar, radius, temp_dir)
                    pic = slide.shapes.add_picture(rp, Inches(bx), Inches(by),
                                                   width=Inches(bw), height=Inches(bh))
                except Exception:
                    pic = None
            if pic is None:
                src = _pptx_photo_path(path, temp_dir)
                pic = slide.shapes.add_picture(src, Inches(bx), Inches(by),
                                               width=Inches(bw), height=Inches(bh))
                try:
                    l, r, t, b = G.crop_fractions(iar, box_ar)
                    pic.crop_left, pic.crop_right = l, r
                    pic.crop_top, pic.crop_bottom = t, b
                except Exception:
                    pass
        elif pic is None:
            src = _pptx_photo_path(path, temp_dir)
            pic = _add_contained_picture(slide, src, bx, by, bw, bh, iar)
        if pic is not None and not baked:
            _set_pic_opacity(pic, opacity)


def _gradient_bg(w, h, c1=(0, 196, 204), c2=(125, 42, 232)):
    base = Image.new("RGB", (w, h), c1)
    top = Image.new("RGB", (w, h), c2)
    mask = Image.new("L", (w, h))
    data = []
    for y in range(h):
        v = int(255 * (y / max(1, h - 1)))
        data.extend([v] * w)
    mask.putdata(data)
    base.paste(top, (0, 0), mask)
    return base


def _dash_font(size):
    return FX._font(size)


def _render_stats_card(stats, temp_dir):
    n = max(1, len(stats))
    TW, TH, GAP = 300, 150, 20
    W, H = n * TW + (n - 1) * GAP, TH
    img = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    f_val, f_lab = _dash_font(44), _dash_font(20)
    for i, (label, value) in enumerate(stats):
        x0 = i * (TW + GAP)
        d.rounded_rectangle([x0, 0, x0 + TW - 1, H - 1], radius=18,
                            fill=(255, 255, 255, 38), outline=(255, 255, 255, 90),
                            width=2)
        vtxt = str(value)
        vb = d.textbbox((0, 0), vtxt, font=f_val)
        d.text((x0 + (TW - (vb[2] - vb[0])) / 2, 24), vtxt, font=f_val,
               fill=(255, 255, 255, 255))
        lb = d.textbbox((0, 0), str(label), font=f_lab)
        d.text((x0 + (TW - (lb[2] - lb[0])) / 2, 96), str(label), font=f_lab,
               fill=(220, 240, 255, 255))
    out = os.path.join(temp_dir, "dash_stats.png")
    img.save(out, "PNG")
    return out


def _render_pie_chart(data_dict, temp_dir):
    W, H = 640, 420
    img = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    colors = [(0, 211, 200), (125, 42, 232), (255, 159, 67), (46, 204, 113),
              (52, 152, 219), (231, 76, 60), (241, 196, 15), (149, 165, 166)]
    items = sorted(data_dict.items(), key=lambda kv: kv[1], reverse=True)
    if len(items) > 7:
        others = sum(v for _, v in items[7:])
        items = items[:7] + [("...", others)]
    total = sum(v for _, v in items) or 1
    box = [20, 20, 400, 400]
    start = -90.0
    font = _dash_font(20)
    for i, (label, val) in enumerate(items):
        extent = 360.0 * val / total
        d.pieslice(box, start, start + extent, fill=colors[i % len(colors)] + (255,),
                   outline=(255, 255, 255, 255), width=2)
        start += extent
    ly = 40
    for i, (label, val) in enumerate(items):
        c = colors[i % len(colors)]
        d.rectangle([430, ly, 454, ly + 24], fill=c + (255,),
                    outline=(255, 255, 255, 200))
        pct = 100.0 * val / total
        d.text((462, ly + 2), "{} — {} ({:.0f}%)".format(str(label)[:14], val, pct),
               font=font, fill=(255, 255, 255, 255))
        ly += 40
    out = os.path.join(temp_dir, "dash_pie.png")
    img.save(out, "PNG")
    return out


def _add_dashboard_slide(prs, layout_slide, job, temp_dir, total_slides):
    """Tương thích cũ — gọi overview đối chiếu ảnh × màn hình."""
    return _add_overview_slides(prs, layout_slide, job, temp_dir)


def _ov_fill_cell(cell, text, *, size=11, bold=False, rgb=(15, 23, 42),
                  fill=None, align='left'):
    cell.text = str(text if text is not None else '')
    try:
        if fill:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(*fill)
        else:
            cell.fill.background()
    except Exception:
        pass
    tf = cell.text_frame
    tf.word_wrap = True
    try:
        tf.margin_left = Pt(4)
        tf.margin_right = Pt(4)
        tf.margin_top = Pt(2)
        tf.margin_bottom = Pt(2)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.name = 'Arial'
    p.font.color.rgb = RGBColor(*rgb)
    p.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT
                   }.get(align, PP_ALIGN.LEFT)


def _ov_status_style(status):
    if status == 'thieu':
        return (254, 226, 226), (153, 27, 27)
    if status == 'du':
        return (220, 252, 231), (22, 101, 52)
    if status == 'thua':
        return (254, 243, 199), (146, 64, 14)
    if status == 'chua_excel':
        return (243, 244, 246), (75, 85, 99)
    return (224, 242, 254), (30, 64, 175)


def _add_overview_slides(prs, layout_slide, job, temp_dir):
    """Trang đầu: từng cửa hàng — số ảnh vs số màn hình (đủ / thiếu)."""
    merged = build_merged_groups(job.excel_by_code)
    ov = build_overview(job.groups, job.excel_by_code, merged)
    rows = ov['rows']
    n_pages = overview_page_count(len(rows))
    added = 0

    def chunk(page):
        if page == 0:
            return rows[:OV_FIRST]
        start = OV_FIRST + (page - 1) * OV_NEXT
        return rows[start:start + OV_NEXT]

    for page in range(n_pages):
        slide = prs.slides.add_slide(layout_slide)
        bg = _gradient_bg(1333, 750)
        bg_path = os.path.join(temp_dir, f"ov_bg_{page}.png")
        bg.save(bg_path, "PNG")
        slide.shapes.add_picture(bg_path, 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)

        tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.22),
                                      Inches(12.4), Inches(0.55))
        p = tb.text_frame.paragraphs[0]
        title = "TỔNG QUAN — ẢNH × MÀN HÌNH"
        if n_pages > 1:
            title += f"   ({page + 1}/{n_pages})"
        p.text = title
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.name = 'Arial'

        sub = slide.shapes.add_textbox(Inches(0.48), Inches(0.72),
                                       Inches(12.3), Inches(0.32))
        sp = sub.text_frame.paragraphs[0]
        extra = ov['excel_no_photo']
        note = ""
        if extra:
            note = f"  ·  Excel còn {len(extra)} điểm chưa có ảnh"
        sp.text = ("Ngày xuất: "
                   + datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
                   + "  ·  Ảnh ≥ màn hình = ĐỦ"
                   + note)
        sp.font.size = Pt(12)
        sp.font.color.rgb = RGBColor(230, 240, 255)
        sp.font.name = 'Arial'

        y_table = 1.12
        if page == 0:
            stats = [("Điểm", ov['n_stores']),
                     ("Ảnh", ov['n_photos']),
                     ("ĐỦ", ov['n_du']),
                     ("THIẾU ảnh", ov['n_thieu'])]
            card_path = _render_stats_card(stats, temp_dir)
            slide.shapes.add_picture(card_path, Inches(0.45), Inches(1.08),
                                     width=Inches(12.4))
            y_table = 2.42

        part = chunk(page)
        n_tbl = len(part) + 1
        headers = ['Mã', 'Tên điểm', 'Ảnh', 'Màn hình', 'Chênh', 'Kết luận']
        widths = [1.9, 4.55, 1.15, 1.45, 1.15, 2.15]
        tbl_h = max(0.7, 0.32 * n_tbl)
        gt = slide.shapes.add_table(
            n_tbl, 6, Inches(0.45), Inches(y_table),
            Inches(12.4), Inches(min(tbl_h, 7.5 - y_table - 0.2)))
        table = gt.table
        table.first_row = True
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
        for i, h in enumerate(headers):
            _ov_fill_cell(table.cell(0, i), h, size=11, bold=True,
                          rgb=(255, 255, 255), fill=(14, 86, 200),
                          align='center' if i >= 2 else 'left')
        for r, rec in enumerate(part, 1):
            bg_c, fg_c = _ov_status_style(rec['status'])
            delta = rec['delta']
            dtxt = f"{delta:+d}" if rec['screens'] else "—"
            vals = [
                rec['code'],
                (rec['name'][:42] + ('…' if len(rec['name']) > 42 else '')),
                rec['photos'],
                rec['screens'] if rec['screens'] else "—",
                dtxt,
                rec['label'],
            ]
            aligns = ['left', 'left', 'center', 'center', 'center', 'center']
            for c, (val, al) in enumerate(zip(vals, aligns)):
                _ov_fill_cell(table.cell(r, c), val, size=11, bold=(c == 5),
                              rgb=fg_c, fill=bg_c, align=al)
        added += 1
    return added


# ════════════════════════════════════════════════════════════
#  Hàm xuất chính
# ════════════════════════════════════════════════════════════
def _try_export_pdf(pptx_path, log):
    """Xuất PDF cạnh PPTX — cần Microsoft PowerPoint trên Windows."""
    pdf_path = os.path.splitext(pptx_path)[0] + '.pdf'
    try:
        import comtypes.client
        pp = comtypes.client.CreateObject('PowerPoint.Application')
        pp.Visible = 1
        pres = pp.Presentations.Open(os.path.abspath(pptx_path), WithWindow=False)
        pres.SaveAs(os.path.abspath(pdf_path), 32)
        pres.Close()
        pp.Quit()
        log(f'PDF: {pdf_path}')
        return pdf_path
    except Exception as e:
        log(f'Không xuất PDF ({os.path.basename(pptx_path)}): {e}')
        return ''


def _export_groups_to_pptx(job, groups, out_path, rep, temp_dir, log, progress_cb,
                         cancel, done_start, total_slides, include_overview=False):
    """Ghi một file PPTX từ dict groups."""
    L = job.layout
    font_cfg = L.get('font', {})
    vis = job.visible or {}
    merged = build_merged_groups(job.excel_by_code)
    fx = job.fx or {}

    def shown(name):
        return vis.get(name, True)

    stem, ext = os.path.splitext(out_path)
    ext = ext or '.pptx'
    prs = _new_prs()
    layout_slide = _blank_layout(prs)
    slide_count, part, done = 0, 1, done_start
    part_cap = G.file_part_limit(job.slides_per_file)

    def save_part(last=False):
        nonlocal prs, layout_slide, slide_count, part
        target = out_path if (last and part == 1) else f"{stem}_Part_{part}{ext}"
        prs.save(target)
        rep.files.append(target)
        log(f'Đã lưu: {target}')
        if job.export_pdf:
            _try_export_pdf(target, log)

    if include_overview and fx.get('dashboard'):
        try:
            n_add = _add_overview_slides(prs, layout_slide, job, temp_dir)
            slide_count += n_add
            rep.slides += n_add
            done += n_add
            log(f'Đã tạo {n_add} slide Overview.')
            if progress_cb:
                progress_cb(done, total_slides)
        except Exception as e:
            log(f'Lỗi overview: {e}')

    sort_hint = ('thứ tự list Excel' if job.sort_mode == 'list'
                 else 'tỉnh/thành Bắc → Nam, rồi quận')
    log(f'Thứ tự slide: {sort_hint}.')

    for code, paths in groups.items():
        if cancel is not None and cancel.is_set():
            raise ExportCancelled()
        paths = sorted(p for p in (paths or []) if p)
        row, mcode, kind = match_row(code, job.excel_by_code, merged)
        if kind == 'fuzzy':
            rep.fuzzy.append((code, mcode))
            log(f"≈ Mã '{code}' khớp gần đúng với Excel '{mcode}'")
        elif kind == 'none' and job.excel_by_code:
            rep.unmatched.append(code)

        name_val = str(clean(row.get('Name'))).strip() or code
        sq = screen_qty(row) if job.pad_blank_slides else 0
        photo_n = len(paths)
        info_rows = build_info_rows(row, photo_n or sq or 1)
        sibs = [r for r in (job.excel_rows or [])
                if str(r.get('Code_RP') or '').strip().upper()
                == str((row or {}).get('Code_RP') or mcode or code or '').upper()]
        info_table = build_info_table(row, sibs or [row])
        nelson_table = (build_nelson_table(row, sibs or [row])
                        if (job.slide_style or 'report') == 'nelson' else None)
        avatar_path = job.avatar_map.get(code.upper())
        if not avatar_path and shown('avatar'):
            rep.missing_avatars.append(code)

        batches = G.build_slide_batches(
            paths, job.n_per_slide, sq, job.pad_blank_slides)
        K = len(batches)
        for k, batch in enumerate(batches, 1):
            if cancel is not None and cancel.is_set():
                raise ExportCancelled()
            if part_cap > 0 and slide_count >= part_cap:
                save_part()
                part += 1
                prs = _new_prs()
                layout_slide = _blank_layout(prs)
                slide_count = 0

            slide = prs.slides.add_slide(layout_slide)
            slide_count += 1
            rep.slides += 1
            if job.bg_image:
                try:
                    _add_bg(slide, prs, job.bg_image)
                except Exception as e:
                    log(f'Lỗi ảnh nền: {e}')

            if shown('image'):
                fx_img = dict(job.fx or {})
                fx_img.update(location_fx_fields(row))
                _place_images(slide, batch, L['image'], job.img_ar, temp_dir, fx_img)

            if shown('avatar') and avatar_path:
                try:
                    _place_avatar(slide, avatar_path, L['avatar'], temp_dir)
                except Exception as e:
                    log(f'Lỗi avatar {code}: {e}')

            if shown('title'):
                tcfg = L['title']
                tx, ty, tw, thh = G.title_box(tcfg)
                tb = slide.shapes.add_textbox(Inches(tx), Inches(ty),
                                              Inches(tw), Inches(thh))
                _style_fit_textbox(
                    tb, G.title_string(tcfg, name_val, k, K), tw, thh,
                    tcfg.get('size', 23),
                    font_name=tcfg.get('font') or font_cfg.get('name', 'Arial'),
                    bold=font_cfg.get('bold', True),
                    italic=font_cfg.get('italic', False),
                    rgb=_hex_rgb(tcfg.get('color') or
                                 font_cfg.get('color') or '#000000'),
                    align=tcfg.get('align', 'left'),
                    max_lines=tcfg.get('max_lines', 2),
                    min_size=tcfg.get('min_size', 12),
                    opacity=tcfg.get('opacity', 100),
                    tracking=tcfg.get('tracking', 0))

            if shown('info'):
                style = job.slide_style or 'report'
                if style == 'saleskit':
                    _add_saleskit_table(slide, info_table, L['info'], font_cfg)
                elif style == 'nelson':
                    _add_nelson_table(slide, nelson_table, L['info'], font_cfg)
                else:
                    _add_info_table(slide, info_rows, L['info'], font_cfg)

            if shown('channel') and job.channel_enabled:
                ccfg = L['channel']
                tb = slide.shapes.add_textbox(
                    Inches(ccfg['x']), Inches(ccfg['y']),
                    Inches(ccfg['w']), Inches(ccfg.get('h', 0.5)))
                ch_txt = (_address_line(row) if (job.slide_style or 'report') == 'nelson'
                          else channel_text(row, job.channel_template))
                _style_fit_textbox(
                    tb, ch_txt,
                    ccfg['w'], ccfg.get('h', 0.5), ccfg.get('size', 10),
                    font_name=ccfg.get('font') or font_cfg.get('name', 'Arial'),
                    bold=font_cfg.get('bold', True),
                    italic=font_cfg.get('italic', False),
                    rgb=_hex_rgb(ccfg.get('color') or
                                 font_cfg.get('color') or '#000000'),
                    align=ccfg.get('align', 'left'),
                    max_lines=ccfg.get('max_lines', 2),
                    min_size=ccfg.get('min_size', 8),
                    opacity=ccfg.get('opacity', 100),
                    tracking=ccfg.get('tracking', 0))

            done += 1
            if progress_cb:
                progress_cb(done, total_slides)

    save_part(last=True)
    return done


def export_pptx(job: ExportJob, progress_cb=None, log_cb=None, cancel=None):
    """Chạy trong worker thread. Trả ExportReport."""
    t0 = time.time()
    rep = ExportReport()
    log = log_cb or (lambda s: None)
    temp_dir = tempfile.mkdtemp(prefix='autopptx2_')

    merged = build_merged_groups(job.excel_by_code)
    groups = filter_groups_by_geo(
        job.groups, job.excel_by_code, merged,
        city=job.city_filter, district=job.district_filter)
    groups = order_groups(
        groups, job.excel_by_code, merged, job.excel_rows, job.sort_mode,
        order_rows=job.sort_rows or job.excel_rows)
    rep.groups = len(groups)

    fx = job.fx or {}
    ov_rows = 0
    if fx.get('dashboard'):
        ov = build_overview(groups, job.excel_by_code, merged)
        ov_rows = len(ov['rows'])
    _, total_slides, _ = estimate(
        groups, job.n_per_slide, overview_rows=ov_rows,
        slides_per_file=job.slides_per_file, by_code=job.excel_by_code,
        merged=merged, pad_blank=job.pad_blank_slides)

    stem, ext = os.path.splitext(job.out_path)
    ext = ext or '.pptx'
    chunks = group_export_chunks(
        groups, job.excel_by_code, merged, job.split_export_by, job.excel_rows)

    try:
        if job.pad_blank_slides:
            log('Bật slide trắng / ô trống theo số màn trên Excel.')
        if job.split_export_by and job.split_export_by != 'none':
            log(f'Tách file theo {job.split_export_by}: {len(chunks)} phần.')
        elif G.file_part_limit(job.slides_per_file):
            log(f'Chia file mỗi {G.file_part_limit(job.slides_per_file)} slide.')
        else:
            log('Gộp mọi slide vào một file PPTX.')

        done = 0
        first = True
        for chunk_name, chunk_groups in chunks.items():
            if not chunk_groups:
                continue
            if chunk_name:
                out_path = f"{stem}_{_safe_filename_part(chunk_name)}{ext}"
            else:
                out_path = job.out_path
            done = _export_groups_to_pptx(
                job, chunk_groups, out_path, rep, temp_dir, log, progress_cb,
                cancel, done, total_slides, include_overview=first and bool(fx.get('dashboard')))
            first = False
    except ExportCancelled:
        rep.cancelled = True
        log('Đã huỷ xuất — không lưu file dở dang.')
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    rep.seconds = time.time() - t0
    return rep


def _save_stamped_image(im, out_path):
    ext = os.path.splitext(out_path)[1].lower()
    if ext in ('.jpg', '.jpeg'):
        im.convert('RGB').save(out_path, 'JPEG', quality=95, subsampling=0)
    elif ext == '.webp':
        im.save(out_path, 'WEBP', quality=95)
    else:
        im.save(out_path, 'PNG')


def export_stamped_images(job: ExportJob, out_dir, progress_cb=None, log_cb=None,
                        cancel=None):
    """Lưu ảnh đã đóng dấu (ngày/vị trí) — không tạo PPTX."""
    t0 = time.time()
    rep = ExportReport(groups=len(job.groups))
    log = log_cb or (lambda s: None)
    os.makedirs(out_dir, exist_ok=True)
    merged = build_merged_groups(job.excel_by_code)
    groups = filter_groups_by_geo(
        job.groups, job.excel_by_code, merged,
        city=job.city_filter, district=job.district_filter)
    groups = order_groups(
        groups, job.excel_by_code, merged, job.excel_rows, job.sort_mode,
        order_rows=job.sort_rows or job.excel_rows)
    items = [(code, p) for code, paths in groups.items() for p in sorted(paths)]
    total = max(1, len(items))
    log(f'Đóng dấu {len(items)} ảnh → {out_dir}')

    for i, (code, path) in enumerate(items):
        if cancel is not None and cancel.is_set():
            rep.cancelled = True
            log('Đã huỷ — giữ ảnh đã lưu.')
            break
        row, mcode, kind = match_row(code, job.excel_by_code, merged)
        if kind == 'fuzzy':
            rep.fuzzy.append((code, mcode))
        elif kind == 'none' and job.excel_by_code:
            rep.unmatched.append(code)

        fx_img = dict(job.fx or {})
        fx_img.update(location_fx_fields(row))
        im = FX.stamp_original(path, fx_img, allow_net=True)
        if im is None:
            log(f'Lỗi ảnh: {os.path.basename(path)}')
            continue

        dest_dir = os.path.join(out_dir, code)
        os.makedirs(dest_dir, exist_ok=True)
        base, ext = os.path.splitext(os.path.basename(path))
        if not ext:
            ext = '.jpg'
        out_path = os.path.join(dest_dir, base + ext)
        if os.path.exists(out_path):
            n = 2
            while os.path.exists(out_path):
                out_path = os.path.join(dest_dir, f'{base}_{n}{ext}')
                n += 1
        try:
            _save_stamped_image(im, out_path)
            rep.files.append(out_path)
            rep.slides += 1
        except Exception as e:
            log(f'Lỗi lưu {os.path.basename(path)}: {e}')
        if progress_cb:
            progress_cb(i + 1, total)

    rep.seconds = time.time() - t0
    return rep
